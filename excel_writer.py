import re
from collections import defaultdict
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Calibri"

# =========================
# STYLE HELPERS
# =========================
def hdr_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def make_font(bold=False, size=11, color="000000"):
    return Font(name=FONT_NAME, bold=bold, size=size, color=color)

def center(wrap=True):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left(wrap=True):
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap)

def right(wrap=False):
    return Alignment(horizontal="right", vertical="center", wrap_text=wrap)

def thin_border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def thick_border():
    t = Side(style="medium")
    return Border(left=t, right=t, top=t, bottom=t)

def style_cell(cell, bold=False, size=11, fg=None, font_color="000000",
               align="left", border="thin"):
    cell.font      = make_font(bold=bold, size=size, color=font_color)
    cell.alignment = center() if align == "center" else left()
    if fg:
        cell.fill = hdr_fill(fg)
    cell.border = thick_border() if border == "thick" else thin_border()

def set_row_height(ws, row, height):
    ws.row_dimensions[row].height = height


# =========================
# MASTER PRODUK & PAKET
# =========================

URUTAN_PRODUK = [
    "Black Garlic 84gr",
    "Black Garlic 100gr",
    "Black Garlic 220gr",
    "Black Garlic 500gr",
    "BG Drink Original",
    "BG Drink Peach",
    "BG Madu Multi Floral",
    "BG Madu Kurma",
    "Box Hampers",
]

# Warna per produk — konsisten di semua sheet
WARNA_PRODUK = {
    "Black Garlic 84gr":      "F4CCCC",
    "Black Garlic 100gr":     "D9EAD3",
    "Black Garlic 220gr":     "CFE2F3",
    "Black Garlic 500gr":     "FCE5CD",
    "BG Drink Original":      "EAD1DC",
    "BG Drink Peach":         "FFF2CC",
    "BG Madu Multi Floral":   "D9D2E9",
    "BG Madu Kurma":          "D0E0E3",
    "Box Hampers":             "F3F3F3",
}

PAKET_KOMPONEN = {
    "Paket 3in1": [
        ("Black Garlic 100gr", 1),
        ("Black Garlic 220gr", 1),
        ("Black Garlic 500gr", 1),
    ],
    "Hampers Imlek": [
        ("Black Garlic 220gr", 2),
        ("Box Hampers", 1),
    ],
    "Hampers Lebaran": [
        ("Black Garlic 220gr", 2),
        ("Box Hampers", 1),
    ],
    "Hampers Natal": [
        ("Black Garlic 220gr", 2),
        ("Box Hampers", 1),
    ],
    "Hampers HSD": [
        ("Black Garlic 220gr", 2),
        ("Box Hampers", 1),
    ],
    "BG Drink Original x2":        [("BG Drink Original", 2)],
    "BG Drink Original x4":        [("BG Drink Original", 4)],
    "BG Drink Original 7 Botol":   [("BG Drink Original", 7)],
    "BG Drink Original x7":        [("BG Drink Original", 7)],
    "BG Drink Peach 7 Botol":      [("BG Drink Peach", 7)],
    "BG Drink Peach x7":           [("BG Drink Peach", 7)],
    "BG Drink Mix 7 Botol": [
        ("BG Drink Peach", 4),
        ("BG Drink Original", 3),
    ],
    "BG Drink Mix 7 Botol (4 Peach + 3 Original)": [
        ("BG Drink Peach", 4),
        ("BG Drink Original", 3),
    ],
    "Black Garlic 100gr x2": [("Black Garlic 100gr", 2)],
    "Black Garlic 100gr x3": [("Black Garlic 100gr", 3)],
    "Black Garlic 100gr x4": [("Black Garlic 100gr", 4)],
    "Black Garlic 100gr x5": [("Black Garlic 100gr", 5)],
    "Black Garlic 220gr x2": [("Black Garlic 220gr", 2)],
    "Black Garlic 220gr x3": [("Black Garlic 220gr", 3)],
    "Black Garlic 220gr x4": [("Black Garlic 220gr", 4)],
    "Black Garlic 220gr x5": [("Black Garlic 220gr", 5)],
    "Black Garlic 500gr x2": [("Black Garlic 500gr", 2)],
    "Black Garlic 500gr x3": [("Black Garlic 500gr", 3)],
    "Black Garlic 500gr x4": [("Black Garlic 500gr", 4)],
    "Black Garlic 500gr x5": [("Black Garlic 500gr", 5)],
    "BG 3in1 (100gr+220gr+500gr)": [
        ("Black Garlic 100gr", 1),
        ("Black Garlic 220gr", 1),
        ("Black Garlic 500gr", 1),
    ],
    "Black Garlic 220gr x2 + Box Hampers Imlek":   [("Black Garlic 220gr", 2), ("Box Hampers", 1)],
    "Black Garlic 220gr x2 + Box Hampers Lebaran": [("Black Garlic 220gr", 2), ("Box Hampers", 1)],
    "Black Garlic 220gr x2 + Box Hampers Natal":   [("Black Garlic 220gr", 2), ("Box Hampers", 1)],
    "Black Garlic 220gr x2 + Box Hampers HSD":     [("Black Garlic 220gr", 2), ("Box Hampers", 1)],
}

PRODUK_TUNGGAL = set(URUTAN_PRODUK)


def expand_to_komponen(nama_produk, qty):
    if not nama_produk:
        return []
    match_x = re.search(r'\s*[xX](\d+)\s*$', nama_produk)
    if match_x:
        multiplier = int(match_x.group(1))
        base_name  = nama_produk[:match_x.start()].strip()
        if base_name in PAKET_KOMPONEN:
            return [(k, q * multiplier * qty) for k, q in PAKET_KOMPONEN[base_name]]
        else:
            return [(base_name, multiplier * qty)]
    if nama_produk in PAKET_KOMPONEN:
        return [(k, q * qty) for k, q in PAKET_KOMPONEN[nama_produk]]
    return [(nama_produk, qty)]


def is_instan(order):
    """Cek apakah ini pesanan instan (Gosend / ojol jemput)."""
    courier  = str(order.get('courier', '')).lower()
    layanan  = str(order.get('layanan', '')).lower()
    platform = str(order.get('platform', '')).lower()
    return (
        'gosend' in courier or 'gosend' in layanan
        or 'instant' in layanan or 'same day' in layanan
        or 'grab' in courier or 'gojek' in courier
    )


# =========================
# PREPARE DATA
# =========================

def prepare(orders):
    """
    Kembalikan semua struktur data yang dibutuhkan semua sheet.

    rekap[akun][shift][nama_komponen]     = total qty botol/pcs
    rekap_resi[akun][shift][nama_komponen] = set of resi (untuk hitung jumlah resi)
    rekap_paket[akun][shift][nama_paket]   = qty paket
    orders_bersih  = semua order valid (termasuk instan)
    orders_instan  = subset dari orders_bersih yang Gosend/ojol
    orders_cek     = order yang tidak terbaca parser
    """
    rekap        = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    rekap_resi   = defaultdict(lambda: defaultdict(lambda: defaultdict(set)))
    rekap_paket  = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    orders_bersih = []
    orders_instan = []
    orders_cek    = []

    for order in orders:
        akun        = order.get('akun', order.get('brand', 'HSD'))
        shift       = order.get('shift', order.get('waktu', 'Pagi'))
        nama_produk = str(order.get('nama_produk', '')).strip()
        qty         = int(order.get('qty', 1) or 1)
        perlu_cek   = order.get('perlu_cek', False)
        resi        = str(order.get('resi', '') or order.get('no_resi', '')).strip()

        if perlu_cek or not nama_produk or nama_produk == '(cek manual)':
            orders_cek.append(order)
            continue

        orders_bersih.append(order)
        if is_instan(order):
            orders_instan.append(order)

        # Hitung paket
        if nama_produk in PAKET_KOMPONEN:
            rekap_paket[akun][shift][nama_produk] += qty
        match_x = re.search(r'\s*[xX](\d+)\s*$', nama_produk)
        if match_x:
            base = nama_produk[:match_x.start()].strip()
            if base in PAKET_KOMPONEN:
                rekap_paket[akun][shift][nama_produk] += qty

        # Expand ke komponen dan rekap qty + resi
        for nama_komponen, qty_komponen in expand_to_komponen(nama_produk, qty):
            if nama_komponen:
                rekap[akun][shift][nama_komponen] += qty_komponen
                if resi:
                    rekap_resi[akun][shift][nama_komponen].add(resi)

    return rekap, rekap_resi, rekap_paket, orders_bersih, orders_instan, orders_cek


# =========================
# SHEET 1 — REKAP GUDANG
# Satu pandang langsung tau: siapkan apa, berapa botol, dari berapa resi
# =========================

def write_sheet1(ws, rekap, rekap_resi, rekap_paket, orders_bersih, orders_cek):
    ws.title = "📦 Rekap Gudang"
    ws.sheet_view.showGridLines = False

    # Lebar kolom
    ws.column_dimensions['A'].width = 32   # Nama produk
    ws.column_dimensions['B'].width = 14   # Qty siapkan
    ws.column_dimensions['C'].width = 14   # Jumlah resi
    ws.column_dimensions['D'].width = 18   # Keterangan

    row = 1

    # ── Banner utama ─────────────────────────────────────────────────────
    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1, value="📦  REKAP KEBUTUHAN GUDANG")
    c.font      = Font(name=FONT_NAME, bold=True, size=16, color="FFFFFF")
    c.fill      = hdr_fill("1F3864")
    c.alignment = center()
    c.border    = thin_border()
    set_row_height(ws, row, 30)
    row += 1

    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1,
                value=f"Dibuat otomatis oleh HSD Agent  ·  {datetime.now().strftime('%d %B %Y  %H:%M WIB')}")
    style_cell(c, size=10, fg="D9E1F2", align="center")
    row += 1

    # ── Ringkasan total ─────────────────────────────────────────────────
    total_resi_unik = len({
        str(o.get('resi', '') or o.get('no_resi', ''))
        for o in orders_bersih
        if str(o.get('resi', '') or o.get('no_resi', ''))
    })
    total_produk = sum(
        sum(data_shift.values())
        for akun_data in rekap.values()
        for data_shift in akun_data.values()
    )

    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1,
                value=f"TOTAL HARI INI:   {total_resi_unik} resi   |   "
                      f"{total_produk} botol/pcs   |   "
                      f"{len(orders_cek)} order perlu cek manual")
    c.font      = Font(name=FONT_NAME, bold=True, size=12, color="000000")
    c.fill      = hdr_fill("FFD966")
    c.alignment = center()
    c.border    = thin_border()
    set_row_height(ws, row, 24)
    row += 1

    # ── Peringatan offline ───────────────────────────────────────────────
    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1,
                value="⚠️  Data ini HANYA dari PDF marketplace — "
                      "order OFFLINE dari Mbak Fitri belum termasuk, tambahkan manual")
    style_cell(c, bold=True, size=10, fg="FFC000", font_color="7F4F00", align="center")
    row += 2

    akun_list  = sorted(rekap.keys())
    shift_list = ["Pagi", "Siang", "Sore"]

    if not akun_list:
        ws.cell(row=row, column=1, value="Tidak ada data order valid.")
        return

    for akun in akun_list:

        # ── Header akun ──────────────────────────────────────────────────
        ws.merge_cells(f"A{row}:D{row}")
        c = ws.cell(row=row, column=1, value=f"🏪  AKUN: {akun}")
        style_cell(c, bold=True, size=13, fg="2E75B6", font_color="FFFFFF", align="center")
        set_row_height(ws, row, 22)
        row += 1

        for shift in shift_list:
            if shift not in rekap[akun]:
                continue

            data_shift  = rekap[akun][shift]
            resi_shift  = rekap_resi[akun].get(shift, {})
            paket_shift = rekap_paket[akun].get(shift, {})

            # Header shift
            ws.merge_cells(f"A{row}:D{row}")
            c = ws.cell(row=row, column=1, value=f"  ⏰  Shift {shift}")
            style_cell(c, bold=True, size=11, fg="BDD7EE", align="left")
            set_row_height(ws, row, 20)
            row += 1

            # Header kolom
            hdrs = ["Nama Produk", "Siapkan (botol/pcs)", "Dari (resi)", "Keterangan"]
            for ci, h in enumerate(hdrs, 1):
                c = ws.cell(row=row, column=ci, value=h)
                style_cell(c, bold=True, size=10, fg="9DC3E6", align="center")
            set_row_height(ws, row, 18)
            row += 1

            # ── Baris paket ──────────────────────────────────────────────
            for nama_paket, qty_paket in sorted(paket_shift.items()):
                c_nama = ws.cell(row=row, column=1, value=f"📦 {nama_paket}")
                c_qty  = ws.cell(row=row, column=2, value=qty_paket)
                c_res  = ws.cell(row=row, column=3, value=f"{qty_paket} paket")
                c_ket  = ws.cell(row=row, column=4, value="")
                style_cell(c_nama, bold=True, size=11, fg="FFF2CC")
                style_cell(c_qty,  bold=True, size=13, fg="FFF2CC", align="center")
                style_cell(c_res,  size=10,   fg="FFF2CC", align="center")
                style_cell(c_ket,  size=10,   fg="FFF2CC")
                set_row_height(ws, row, 20)
                row += 1

                # Komponen paket
                match_x   = re.search(r'\s*[xX](\d+)\s*$', nama_paket)
                base_name  = nama_paket[:match_x.start()].strip() if match_x else nama_paket
                multiplier = int(match_x.group(1)) if match_x else 1
                if base_name in PAKET_KOMPONEN:
                    for komponen, qty_per_unit in PAKET_KOMPONEN[base_name]:
                        total_k = qty_per_unit * multiplier * qty_paket
                        c_k = ws.cell(row=row, column=1, value=f"     → {komponen}")
                        c_q = ws.cell(row=row, column=2, value=total_k)
                        c_s = ws.cell(row=row, column=3, value="")
                        c_e = ws.cell(row=row, column=4, value=f"isi {qty_per_unit * multiplier}/paket")
                        style_cell(c_k, size=10, fg="FFFACD")
                        style_cell(c_q, bold=True, size=11, fg="FFFACD", align="center")
                        style_cell(c_s, size=10, fg="FFFACD", align="center")
                        style_cell(c_e, size=10, fg="FFFACD")
                        row += 1

            # ── Baris produk tunggal ─────────────────────────────────────
            for nama_produk in URUTAN_PRODUK:
                if nama_produk not in data_shift:
                    continue
                qty_total   = data_shift[nama_produk]
                jumlah_resi = len(resi_shift.get(nama_produk, set()))
                bg          = WARNA_PRODUK.get(nama_produk, "F2F2F2")

                c_nama = ws.cell(row=row, column=1, value=nama_produk)
                c_qty  = ws.cell(row=row, column=2, value=qty_total)
                c_res  = ws.cell(row=row, column=3,
                                 value=jumlah_resi if jumlah_resi > 0 else "")
                c_ket  = ws.cell(row=row, column=4, value="")

                style_cell(c_nama, size=11,   fg=bg)
                style_cell(c_qty,  bold=True, size=14, fg=bg, align="center")
                style_cell(c_res,  bold=True, size=12, fg=bg, align="center")
                style_cell(c_ket,  size=10,   fg=bg)
                set_row_height(ws, row, 22)
                row += 1

            # Produk yang tidak ada di URUTAN_PRODUK (fallback)
            for nama_produk, qty_total in data_shift.items():
                if nama_produk in PRODUK_TUNGGAL:
                    continue
                jumlah_resi = len(resi_shift.get(nama_produk, set()))
                c_nama = ws.cell(row=row, column=1, value=nama_produk)
                c_qty  = ws.cell(row=row, column=2, value=qty_total)
                c_res  = ws.cell(row=row, column=3,
                                 value=jumlah_resi if jumlah_resi > 0 else "")
                c_ket  = ws.cell(row=row, column=4, value="⚠️ cek mapping")
                style_cell(c_nama, size=10, fg="FFE699")
                style_cell(c_qty,  bold=True, size=12, fg="FFE699", align="center")
                style_cell(c_res,  size=10, fg="FFE699", align="center")
                style_cell(c_ket,  size=10, fg="FFE699")
                row += 1

            # Total shift
            total_qty_shift  = sum(data_shift.values())
            total_resi_shift = len({
                r for prod_set in resi_shift.values() for r in prod_set
            })
            c_t  = ws.cell(row=row, column=1,
                           value=f"  ✅  TOTAL SHIFT {shift.upper()}")
            c_tq = ws.cell(row=row, column=2, value=total_qty_shift)
            c_tr = ws.cell(row=row, column=3, value=total_resi_shift)
            c_te = ws.cell(row=row, column=4, value="")
            style_cell(c_t,  bold=True, size=11, fg="70AD47", font_color="FFFFFF")
            style_cell(c_tq, bold=True, size=14, fg="70AD47", font_color="FFFFFF",
                       align="center")
            style_cell(c_tr, bold=True, size=13, fg="70AD47", font_color="FFFFFF",
                       align="center")
            style_cell(c_te, fg="70AD47")
            set_row_height(ws, row, 22)
            row += 2

        # Total keseluruhan per akun
        total_akun      = defaultdict(int)
        total_resi_akun = defaultdict(set)
        for shift in shift_list:
            for nama, qty in rekap[akun].get(shift, {}).items():
                total_akun[nama] += qty
            for nama, resi_set in rekap_resi[akun].get(shift, {}).items():
                total_resi_akun[nama].update(resi_set)

        ws.merge_cells(f"A{row}:D{row}")
        c_ta = ws.cell(row=row, column=1,
                       value=f"TOTAL KESELURUHAN  —  {akun}")
        style_cell(c_ta, bold=True, size=12, fg="1F3864", font_color="FFFFFF",
                   align="center")
        set_row_height(ws, row, 22)
        row += 1

        grand_qty  = 0
        grand_resi = set()
        for nama_produk in URUTAN_PRODUK:
            if nama_produk not in total_akun:
                continue
            q = total_akun[nama_produk]
            r = len(total_resi_akun.get(nama_produk, set()))
            grand_qty  += q
            grand_resi.update(total_resi_akun.get(nama_produk, set()))
            bg = WARNA_PRODUK.get(nama_produk, "D6DCE4")
            c_n = ws.cell(row=row, column=1, value=nama_produk)
            c_q = ws.cell(row=row, column=2, value=q)
            c_r = ws.cell(row=row, column=3, value=r if r > 0 else "")
            c_e = ws.cell(row=row, column=4, value="")
            style_cell(c_n, bold=True, size=11, fg=bg)
            style_cell(c_q, bold=True, size=14, fg=bg, align="center")
            style_cell(c_r, bold=True, size=12, fg=bg, align="center")
            style_cell(c_e, fg=bg)
            set_row_height(ws, row, 22)
            row += 1

        c_gn = ws.cell(row=row, column=1, value=f"  GRAND TOTAL — {akun}")
        c_gq = ws.cell(row=row, column=2, value=grand_qty)
        c_gr = ws.cell(row=row, column=3, value=len(grand_resi))
        c_ge = ws.cell(row=row, column=4, value="")
        style_cell(c_gn, bold=True, size=12, fg="1F3864", font_color="FFFFFF")
        style_cell(c_gq, bold=True, size=16, fg="1F3864", font_color="FFFF00",
                   align="center")
        style_cell(c_gr, bold=True, size=14, fg="1F3864", font_color="FFFF00",
                   align="center")
        style_cell(c_ge, fg="1F3864")
        set_row_height(ws, row, 26)
        row += 3

    # Keterangan offline di bagian bawah
    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1, value="📋  CATATAN PENTING")
    style_cell(c, bold=True, size=11, fg="FFC000", font_color="000000", align="center")
    row += 1
    for baris in [
        "• Data di atas HANYA dari PDF marketplace (Shopee, TikTok, Lazada, dll).",
        "• Order OFFLINE dari Mbak Fitri TIDAK masuk — tambahkan manual.",
        "• Kolom 'Dari (resi)' = berapa nomor resi berbeda yang pesan produk itu.",
        "• Kolom 'Siapkan' = total botol/pcs yang harus disiapkan.",
    ]:
        ws.merge_cells(f"A{row}:D{row}")
        c = ws.cell(row=row, column=1, value=baris)
        style_cell(c, size=10, fg="FFF2CC", font_color="7F6000")
        row += 1


# =========================
# SHEET 2 — DETAIL SEMUA RESI (untuk Viora centang saat print)
# =========================

def write_sheet2_detail_resi(ws, orders_bersih):
    ws.title = "📋 Detail Semua Resi"
    ws.sheet_view.showGridLines = False

    cols   = ["No", "No. Resi", "Platform", "Kurir", "Produk", "Qty",
              "Akun", "Shift", "Pembayaran", "Penerima"]
    widths = [5, 26, 12, 14, 30, 6, 14, 8, 12, 25]

    for i, (col, w) in enumerate(zip(cols, widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1

    ws.merge_cells(f"A{row}:J{row}")
    c = ws.cell(row=row, column=1, value="📋  DETAIL SEMUA RESI  —  untuk Viora (centang saat print)")
    c.font      = Font(name=FONT_NAME, bold=True, size=13, color="FFFFFF")
    c.fill      = hdr_fill("1F3864")
    c.alignment = center()
    c.border    = thin_border()
    set_row_height(ws, row, 26)
    row += 1

    ws.merge_cells(f"A{row}:J{row}")
    c = ws.cell(row=row, column=1,
                value=f"Total: {len(orders_bersih)} order valid  ·  "
                      f"{datetime.now().strftime('%d %B %Y  %H:%M WIB')}")
    style_cell(c, size=10, fg="D9E1F2", align="center")
    row += 1

    # Header kolom
    for ci, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=ci, value=h)
        style_cell(c, bold=True, size=10, fg="2E75B6", font_color="FFFFFF", align="center")
    set_row_height(ws, row, 18)
    row += 1

    # Sort: akun → shift → platform → resi
    shift_order = {"Pagi": 0, "Siang": 1, "Sore": 2}
    sorted_orders = sorted(
        orders_bersih,
        key=lambda o: (
            o.get('akun', ''),
            shift_order.get(o.get('shift', 'Pagi'), 9),
            o.get('platform', ''),
            o.get('resi', '') or o.get('no_resi', ''),
        )
    )

    colors = ["FFFFFF", "F2F7FF"]
    current_akun_shift = None

    for i, order in enumerate(sorted_orders):
        akun  = order.get('akun', '')
        shift = order.get('shift', '')
        key   = (akun, shift)

        # Sisipkan pemisah akun/shift
        if key != current_akun_shift:
            current_akun_shift = key
            ws.merge_cells(f"A{row}:J{row}")
            c = ws.cell(row=row, column=1,
                        value=f"▶  {akun}  —  Shift {shift}")
            style_cell(c, bold=True, size=11, fg="BDD7EE", align="left")
            set_row_height(ws, row, 20)
            row += 1

        bg = colors[i % 2]
        resi = order.get('resi', '') or order.get('no_resi', '')
        vals = [
            i + 1,
            resi,
            order.get('platform', ''),
            order.get('courier', ''),
            order.get('nama_produk', ''),
            order.get('qty', 1),
            akun,
            shift,
            order.get('pembayaran', ''),
            order.get('nama_penerima', ''),
        ]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=ci, value=val)
            if ci == 5:  # kolom produk
                produk_bg = WARNA_PRODUK.get(order.get('nama_produk', ''), bg)
                style_cell(c, bold=True, size=10, fg=produk_bg)
            elif ci == 6:  # kolom qty
                style_cell(c, bold=True, size=12, fg=bg, align="center")
            elif ci == 2:  # kolom resi
                style_cell(c, size=10, fg=bg)
            else:
                style_cell(c, size=10, fg=bg)
        set_row_height(ws, row, 18)
        row += 1


# =========================
# SHEET 3 — PESANAN INSTAN (Gosend / ojol jemput)
# =========================

def write_sheet3_instan(ws, orders_instan):
    ws.title = "⚡ Pesanan Instan"
    ws.sheet_view.showGridLines = False

    cols   = ["No", "No. Resi", "Platform", "Kurir", "Produk", "Qty",
              "Akun", "Shift", "Pembayaran", "Penerima"]
    widths = [5, 26, 12, 14, 30, 6, 14, 8, 12, 25]

    for i, (col, w) in enumerate(zip(cols, widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1

    ws.merge_cells(f"A{row}:J{row}")
    c = ws.cell(row=row, column=1,
                value="⚡  PESANAN INSTAN  —  Gosend / Ojol Jemput  —  SIAPKAN DULUAN!")
    c.font      = Font(name=FONT_NAME, bold=True, size=13, color="FFFFFF")
    c.fill      = hdr_fill("C00000")
    c.alignment = center()
    c.border    = thin_border()
    set_row_height(ws, row, 26)
    row += 1

    if not orders_instan:
        ws.merge_cells(f"A{row}:J{row}")
        c = ws.cell(row=row, column=1, value="Tidak ada pesanan instan hari ini.")
        style_cell(c, size=11, fg="F2F2F2", align="center")
        return

    ws.merge_cells(f"A{row}:J{row}")
    c = ws.cell(row=row, column=1,
                value=f"Total: {len(orders_instan)} pesanan instan  ·  "
                      f"{datetime.now().strftime('%d %B %Y  %H:%M WIB')}")
    style_cell(c, size=10, fg="FCE4D6", align="center")
    row += 1

    for ci, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=ci, value=h)
        style_cell(c, bold=True, size=10, fg="C00000", font_color="FFFFFF", align="center")
    set_row_height(ws, row, 18)
    row += 1

    for i, order in enumerate(orders_instan):
        bg   = "FFF2CC" if i % 2 == 0 else "FFFDE7"
        resi = order.get('resi', '') or order.get('no_resi', '')
        vals = [
            i + 1, resi,
            order.get('platform', ''),
            order.get('courier', ''),
            order.get('nama_produk', ''),
            order.get('qty', 1),
            order.get('akun', ''),
            order.get('shift', ''),
            order.get('pembayaran', ''),
            order.get('nama_penerima', ''),
        ]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=ci, value=val)
            style_cell(c, bold=(ci == 6), size=11 if ci == 6 else 10,
                       fg=bg, align="center" if ci in (1, 6) else "left")
        set_row_height(ws, row, 18)
        row += 1


# =========================
# SHEET 4 — VARIASI ORDER PER RESI
# =========================

def write_sheet4_variasi(ws, orders_bersih):
    ws.title = "📊 Variasi Order"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 45
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 28

    row = 1

    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1, value="📊  VARIASI ORDER PER RESI")
    c.font      = Font(name=FONT_NAME, bold=True, size=13, color="FFFFFF")
    c.fill      = hdr_fill("1F3864")
    c.alignment = center()
    c.border    = thin_border()
    set_row_height(ws, row, 26)
    row += 1

    ws.merge_cells(f"A{row}:D{row}")
    c = ws.cell(row=row, column=1,
                value=f"Generated: {datetime.now().strftime('%d %B %Y %H:%M WIB')}")
    style_cell(c, size=10, fg="D9E1F2", align="center")
    row += 2

    variasi     = defaultdict(list)
    resi_produk = defaultdict(list)

    for order in orders_bersih:
        resi = order.get('resi', '') or order.get('no_resi', '')
        nama = order.get('nama_produk', '')
        qty  = int(order.get('qty', 1) or 1)
        if resi and nama:
            resi_produk[resi].append((nama, qty))

    for resi, produk_list in resi_produk.items():
        key = tuple(sorted([(n, q) for n, q in produk_list]))
        variasi[key].append(resi)

    variasi_sorted = sorted(variasi.items(), key=lambda x: -len(x[1]))

    headers = ["Kombinasi Produk yang Dibeli", "Jumlah Resi", "% dari Total", "Contoh Resi"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=ci, value=h)
        style_cell(c, bold=True, size=10, fg="2E75B6", font_color="FFFFFF", align="center")
    row += 1

    total_resi = len(resi_produk)
    colors = ["DDEEFF", "EEF5FF"]

    for i, (key, resi_list) in enumerate(variasi_sorted):
        bg = colors[i % 2]
        kombinasi = (
            f"{key[0][0]} × {key[0][1]}" if len(key) == 1
            else " | ".join(f"{n} × {q}" for n, q in key)
        )
        jumlah = len(resi_list)
        pct    = f"{jumlah / total_resi * 100:.1f}%" if total_resi > 0 else "0%"
        contoh = resi_list[0] if resi_list else ""

        c1 = ws.cell(row=row, column=1, value=kombinasi)
        c2 = ws.cell(row=row, column=2, value=jumlah)
        c3 = ws.cell(row=row, column=3, value=pct)
        c4 = ws.cell(row=row, column=4, value=contoh)
        style_cell(c1, size=10, fg=bg)
        style_cell(c2, bold=True, size=11, fg=bg, align="center")
        style_cell(c3, size=10, fg=bg, align="center")
        style_cell(c4, size=10, fg=bg)
        row += 1

    row += 1
    c_t = ws.cell(row=row, column=1, value="TOTAL RESI")
    c_v = ws.cell(row=row, column=2, value=total_resi)
    style_cell(c_t, bold=True, size=11, fg="1F3864", font_color="FFFFFF")
    style_cell(c_v, bold=True, size=13, fg="1F3864", font_color="FFFFFF", align="center")
    for col in [3, 4]:
        style_cell(ws.cell(row=row, column=col), fg="1F3864", font_color="FFFFFF")


# =========================
# SHEET 5 — PERLU TINDAKAN MANUAL
# =========================

def write_sheet5_perlu_cek(ws, orders_cek):
    ws.title = "⚠️ Perlu Tindakan"
    ws.sheet_view.showGridLines = False

    cols   = ["Resi", "Platform", "SKU Raw", "Nama Produk", "Qty",
              "Alasan", "Akun", "Shift"]
    widths = [26, 12, 36, 30, 8, 32, 12, 10]

    for i, (col, w) in enumerate(zip(cols, widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1
    ws.merge_cells(f"A{row}:H{row}")
    c = ws.cell(row=row, column=1,
                value=f"⚠️  ORDER PERLU CEK MANUAL  —  {len(orders_cek)} order")
    c.font      = Font(name=FONT_NAME, bold=True, size=12, color="FFFFFF")
    c.fill      = hdr_fill("C00000")
    c.alignment = center()
    c.border    = thin_border()
    set_row_height(ws, row, 24)
    row += 1

    ws.merge_cells(f"A{row}:H{row}")
    c = ws.cell(row=row, column=1,
                value=f"Total: {len(orders_cek)} order  ·  "
                      f"{datetime.now().strftime('%d %B %Y %H:%M WIB')}")
    style_cell(c, size=10, fg="FCE4D6", align="center")
    row += 1

    for ci, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=ci, value=h)
        style_cell(c, bold=True, size=10, fg="C00000", font_color="FFFFFF", align="center")
    row += 1

    for i, order in enumerate(orders_cek):
        bg = "FFF2CC" if i % 2 == 0 else "FFFFFF"
        vals = [
            order.get('resi', '') or order.get('no_resi', ''),
            order.get('platform', ''),
            order.get('sku_raw', '') or order.get('sku', ''),
            order.get('nama_produk', ''),
            order.get('qty', ''),
            order.get('alasan_cek', 'SKU tidak dikenali'),
            order.get('akun', ''),
            order.get('shift', ''),
        ]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=ci, value=val)
            style_cell(c, size=10, fg=bg)
        row += 1


# =========================
# MAIN WRITE FUNCTION
# =========================

def write_excel(orders, output_path):
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    rekap, rekap_resi, rekap_paket, orders_bersih, orders_instan, orders_cek = prepare(orders)

    # Sheet 1 — Rekap Gudang (utama, yang dibuka pertama)
    ws1 = wb.create_sheet("📦 Rekap Gudang")
    write_sheet1(ws1, rekap, rekap_resi, rekap_paket, orders_bersih, orders_cek)

    # Sheet 2 — Detail semua resi (untuk Viora centang)
    ws2 = wb.create_sheet("📋 Detail Semua Resi")
    write_sheet2_detail_resi(ws2, orders_bersih)

    # Sheet 3 — Pesanan instan (Gosend/ojol, siapkan duluan)
    ws3 = wb.create_sheet("⚡ Pesanan Instan")
    write_sheet3_instan(ws3, orders_instan)

    # Sheet 4 — Variasi order (analisis kombinasi)
    ws4 = wb.create_sheet("📊 Variasi Order")
    write_sheet4_variasi(ws4, orders_bersih)

    # Sheet 5 — Perlu tindakan (hanya jika ada)
    if orders_cek:
        ws5 = wb.create_sheet("⚠️ Perlu Tindakan")
        write_sheet5_perlu_cek(ws5, orders_cek)

    wb.save(output_path)
    return {
        "total_order":   len(orders),
        "order_valid":   len(orders_bersih),
        "order_instan":  len(orders_instan),
        "order_cek":     len(orders_cek),
        "output_path":   output_path,
    }


# =========================
# TEST / DEBUG
# =========================

if __name__ == "__main__":
    test_orders = [
        {"akun": "HSD Jakarta", "shift": "Pagi",  "resi": "TK001",
         "platform": "TikTok",  "courier": "J&T",   "layanan": "REG",
         "nama_produk": "Black Garlic 100gr",  "qty": 1, "perlu_cek": False,
         "pembayaran": "Non-COD", "nama_penerima": "Budi Santoso"},
        {"akun": "HSD Jakarta", "shift": "Pagi",  "resi": "TK002",
         "platform": "TikTok",  "courier": "J&T",   "layanan": "REG",
         "nama_produk": "Black Garlic 220gr",  "qty": 2, "perlu_cek": False,
         "pembayaran": "Non-COD", "nama_penerima": "Siti Rahayu"},
        {"akun": "HSD Jakarta", "shift": "Pagi",  "resi": "TK003",
         "platform": "Shopee",  "courier": "SiCepat", "layanan": "REG",
         "nama_produk": "Black Garlic 500gr",  "qty": 1, "perlu_cek": False,
         "pembayaran": "COD", "nama_penerima": "Ahmad Fauzi"},
        {"akun": "HSD Jakarta", "shift": "Pagi",  "resi": "TK004",
         "platform": "TikTok",  "courier": "J&T",   "layanan": "REG",
         "nama_produk": "Paket 3in1",           "qty": 2, "perlu_cek": False,
         "pembayaran": "Non-COD", "nama_penerima": "Dewi Lestari"},
        {"akun": "HSD Jakarta", "shift": "Pagi",  "resi": "GO001",
         "platform": "Shopee",  "courier": "Gosend", "layanan": "Instant",
         "nama_produk": "Black Garlic 220gr",  "qty": 1, "perlu_cek": False,
         "pembayaran": "Non-COD", "nama_penerima": "Rina Wati"},
        {"akun": "HSD Jakarta", "shift": "Siang", "resi": "SP001",
         "platform": "Shopee",  "courier": "SiCepat", "layanan": "REG",
         "nama_produk": "Black Garlic 100gr",  "qty": 3, "perlu_cek": False,
         "pembayaran": "COD", "nama_penerima": "Hendra Kusuma"},
        {"akun": "HSS",         "shift": "Pagi",  "resi": "LZ001",
         "platform": "Lazada",  "courier": "LEX",   "layanan": "STANDARD",
         "nama_produk": "Black Garlic 220gr",  "qty": 1, "perlu_cek": False,
         "pembayaran": "Non-COD", "nama_penerima": "Maya Sari"},
        {"akun": "HSS",         "shift": "Pagi",  "resi": "LZ002",
         "platform": "Lazada",  "courier": "LEX",   "layanan": "STANDARD",
         "nama_produk": "BG Drink Mix 7 Botol","qty": 1, "perlu_cek": False,
         "pembayaran": "Non-COD", "nama_penerima": "Fajar Nugroho"},
        {"akun": "HSD Jakarta", "shift": "Pagi",  "resi": "XX001",
         "platform": "Unknown", "courier": "Unknown", "layanan": "-",
         "nama_produk": "(cek manual)", "qty": 1, "perlu_cek": True,
         "sku_raw": "BG-UNKNOWN-SKU", "alasan_cek": "SKU tidak dikenali",
         "nama_penerima": ""},
    ]

    out = "/tmp/test_excel_writer.xlsx"
    result = write_excel(test_orders, out)
    print(f"✅ Excel berhasil: {out}")
    print(f"   Total: {result['total_order']} | Valid: {result['order_valid']} "
          f"| Instan: {result['order_instan']} | Cek: {result['order_cek']}")
